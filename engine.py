"""Export parsers and explainable monitoring rules. No writes to source systems."""
import base64, csv, hashlib, io, json, math, sqlite3
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

KINDS = ['inventory','custom_orders','bulk_orders','po_invoices','ap_summary','eam','manual','departures']
def day(v):
    if not v: return None
    try: return datetime.fromisoformat(str(v).replace('Z','+00:00')).date()
    except ValueError:
        for fmt in ('%d-%b-%Y','%d-%b-%y','%m/%d/%Y'):
            try: return datetime.strptime(str(v),fmt).date()
            except ValueError: pass
    return None
def num(v):
    try: return float(v or 0)
    except (ValueError,TypeError): return 0
def serial(v): return str(v or '').strip().upper()
def clean(v): return v.isoformat(sep=' ') if isinstance(v,datetime) else v.isoformat() if isinstance(v,date) else v
def reference(import_id, filename, sheet, row, values):
    return dict(import_id=import_id,file=filename,sheet=sheet,row=row,values=values)
def parse(content, filename, kind):
    if kind not in KINDS: raise ValueError('Select a supported source type.')
    if filename.lower().endswith('.csv'):
        raw=list(csv.reader(io.StringIO(content.decode('utf-8-sig')))); sheet='CSV'
    elif filename.lower().endswith('.xlsx'):
        import openpyxl
        wb=openpyxl.load_workbook(io.BytesIO(content),data_only=True,read_only=True)
        if len(wb.sheetnames)!=1: raise ValueError('Upload a single-sheet source export; reconciliation workbooks are reference outputs.')
        ws=wb.active;sheet=ws.title;raw=[[clean(v) for v in r] for r in ws.iter_rows(values_only=True)];wb.close()
    else: raise ValueError('Supported files: .xlsx and UTF-8 .csv.')
    records=[]
    if kind=='eam':
        section=None
        for rn,r in enumerate(raw,1):
            first=str(r[0] or '').strip() if r else ''
            if first in ['Success Line Details','Error Line Details','Warning Line Details']:
                section=first.split()[0].lower();continue
            if section and len(r)>7 and r[1] is not None and str(r[7] or '').strip() not in ('','Asset Tag'):
                get=lambda i:r[i] if len(r)>i else None
                records.append(dict(serial=serial(get(7)),billing_id=str(get(1)),po=str(get(3) or ''),assigned=get(4),request=get(5),ritm=get(6),period=get(8),amount=num(get(14)),asset_amount=num(get(11)),status=section,message=str(get(19) or '').strip(),account=get(15),row=rn,cells={f'{column(i+1)}{rn}':v for i,v in enumerate(r) if v is not None}))
        if not any(str(r[0] or '').strip()=='Success Line Details' for r in raw if r): raise ValueError('EAM format not recognized: missing Success Line Details.')
    else:
        if not raw: raise ValueError('File has no rows.')
        headers=[str(v or '').strip() for v in raw[0]]
        required={'inventory':['serial_number','model','install_status'],'bulk_orders':['purchase_order','ordered_quantity','status'],'custom_orders':['purchase_order','ordered_quantity','status'],'po_invoices':['Purchase Order Number','P2P Transaction Type Description','AP Invoice Number'],'ap_summary':['AP Invoice Number','AP Invoice Amount'],'manual':['ASSET_NUMBER','TRANSACTION_TYPE','PRICE'],'departures':['person_id','person_name','departure_date']}[kind]
        missing=[k for k in required if k not in headers]
        if missing: raise ValueError('Missing fields: '+', '.join(missing))
        for rn,r in enumerate(raw[1:],2):
            if not any(v is not None and str(v).strip() for v in r): continue
            vals=dict(zip(headers,r)); vals.pop('',None)
            if kind=='departures' and (not vals.get('person_id') or not day(vals.get('departure_date'))): raise ValueError(f'Row {rn}: person_id and an ISO departure_date are required.')
            records.append(dict(values=vals,row=rn,cells={f'{column(i+1)}{rn}':v for i,v in enumerate(r) if v is not None}))
    if not records: raise ValueError('No data records found.')
    return sheet,records
def column(i):
    s=''
    while i:i,n=divmod(i-1,26);s=chr(65+n)+s
    return s
def connect(path):
    db=sqlite3.connect(path,timeout=20);db.row_factory=sqlite3.Row
    db.executescript('''CREATE TABLE IF NOT EXISTS imports(id INTEGER PRIMARY KEY,kind TEXT,name TEXT,asof TEXT,sha TEXT,sheet TEXT,records TEXT,created TEXT,UNIQUE(kind,asof,sha));
    CREATE TABLE IF NOT EXISTS decisions(id TEXT PRIMARY KEY,state TEXT,owner TEXT,note TEXT,signature TEXT,updated TEXT);
    CREATE TABLE IF NOT EXISTS settings(id INTEGER PRIMARY KEY,value TEXT);''')
    return db
def ingest(db,content,filename,kind,asof):
    parsed_day=day(asof)
    if not parsed_day: raise ValueError('A valid export date is required.')
    digest=hashlib.sha256(content).hexdigest()
    old=db.execute('SELECT id FROM imports WHERE kind=? AND asof=? AND sha=?',(kind,parsed_day.isoformat(),digest)).fetchone()
    if old:return dict(id=old['id'],duplicate=True)
    sheet,records=parse(content,filename,kind)
    cur=db.execute('INSERT INTO imports(kind,name,asof,sha,sheet,records,created) VALUES(?,?,?,?,?,?,?)',(kind,Path(filename).name,parsed_day.isoformat(),digest,sheet,json.dumps(records),datetime.now().isoformat()))
    db.commit();return dict(id=cur.lastrowid,duplicate=False,rows=len(records))
def imports(db):
    return [dict(r) for r in db.execute('SELECT * FROM imports ORDER BY asof,id')]
def unpack(imp):
    return [dict(r,source=reference(imp['id'],imp['name'],imp['sheet'],r['row'],r.get('cells',{})),asof=imp['asof']) for r in json.loads(imp['records'])]
def latest(all_imports,kind):
    its=[i for i in all_imports if i['kind']==kind]
    return unpack(its[-1]) if its else []
def index_assets(records):
    assets={};duplicates=[]
    for r in records:
        v=r['values'];key=serial(v.get('serial_number'))
        if not key:continue
        if key in assets:duplicates.append(key)
        assets[key]=dict(serial=key,model=v.get('model') or 'Unknown model',location=v.get('purchase_line.purchase_order.ship_to') or 'Unknown ship-to',state=v.get('install_status'),assignee=v.get('assigned_to'),assigned=v.get('assigned'),billed=str(v.get('u_billed')).lower()=='true',cost=num(v.get('cost')),base_cost=num(v.get('u_base_cost')),po=str(v.get('po_number') or ''),ticket=v.get('purchase_line.purchase_order.number'),received=v.get('purchase_line.received'),person_id=v.get('assigned_to.user_name') or v.get('assigned_to.sys_id'),source=r['source'],asof=r['asof'])
    return assets,duplicates
def signature(issue):
    return hashlib.sha256(json.dumps([issue['type'],issue['serial'],issue.get('amount'),issue['description']],sort_keys=True).encode()).hexdigest()
def build(db):
    imps=imports(db);inv=[i for i in imps if i['kind']=='inventory'];inventory_rows=latest(imps,'inventory');assets,dupes=index_assets(inventory_rows)
    previous,_=index_assets(unpack(inv[-2])) if len(inv)>1 else ({},[])
    asof=day(inv[-1]['asof']) if inv else date.today();prevday=day(inv[-2]['asof']) if len(inv)>1 else asof
    span=max(1,(asof-prevday).days)
    settings=dict(lead_days=21,buffer_days=7,grace_days=30)
    saved=db.execute('SELECT value FROM settings WHERE id=1').fetchone()
    if saved:settings.update(json.loads(saved['value']))
    # Latest order snapshots only: historical canceled/received orders are not inbound supply.
    orders=[]
    for kind in ('bulk_orders','custom_orders'):
        for r in latest(imps,kind):
            v=r['values'];orders.append(dict(ticket=v.get('purchase_order'),po=str(v.get('purchase_order.u_oracle_po') or ''),model=v.get('model'),location=v.get('purchase_order.ship_to') or 'Unknown ship-to',quantity=num(v.get('ordered_quantity')),status=v.get('status'),eta=v.get('expected_delivery'),source=r['source'],asof=r['asof']))
    events=[];seen=set()
    for imp in imps:
        if imp['kind'] not in ('eam','manual'):continue
        for r in unpack(imp):
            if imp['kind']=='eam':
                e=dict(serial=r['serial'],status=r['status'],amount=r['amount'],request=r.get('request'),ritm=r.get('ritm'),period=r.get('period'),message=r['message'],account=r.get('account'),event_id=r['billing_id'],source=r['source'],asof=r['asof'],type='eam')
            else:
                v=r['values'];e=dict(serial=serial(v.get('ASSET_NUMBER')),status='prepared',amount=num(v.get('QUANTITY'))*num(v.get('PRICE'))*(1 if v.get('TRANSACTION_TYPE')=='CHARGE' else -1),request=None,ritm=None,period=v.get('BILLING_PERIOD'),message=f"{v.get('TRANSACTION_TYPE')} input · {v.get('TRANSACTION_LINE_DESCRIPTION')}",event_id=f"{v.get('SOURCE_BATCH_NUMBER')}:{v.get('SOURCE_TRANSACTION_ID')}",source=r['source'],asof=r['asof'],type='manual')
            key=(e['type'],e['serial'],e['event_id'],e['status'],e['period'])
            if key not in seen:seen.add(key);events.append(e)
    byasset=defaultdict(list)
    for e in events:byasset[e['serial']].append(e)
    finance=defaultdict(list)
    for r in latest(imps,'po_invoices'):
        finance[str(r['values'].get('Purchase Order Number') or '')].append(r)
    issues=[]
    def add(kind,key,asset,title,desc,amount=None,evidence=None,priority='Review',action=''):
        a=asset or {};issue=dict(id=kind+':'+key,type=kind,serial=a.get('serial',key),title=title,description=desc,amount=amount,model=a.get('model','Asset not in latest inventory'),location=a.get('location','Unknown'),assignee=a.get('assignee'),priority=priority,action=action,evidence=evidence or [],state='Open',owner='',note='')
        issue['signature']=signature(issue);issues.append(issue)
    for e in events:
        if e['status']!='error':continue
        success=[x for x in byasset[e['serial']] if x['status']=='success' and x['asof']>=e['asof'] and x.get('request')==e.get('request') and x.get('ritm')==e.get('ritm') and abs(x['amount']-e['amount'])<=1.5]
        if success:continue
        a=assets.get(e['serial'],dict(serial=e['serial']));add('rejected',e['serial']+':'+e['event_id'],a,'Billing rejected',e['message']+f" Attempted in export dated {e['asof']}; no matching later EAM success found.",e['amount'],[e['source']],priority='High',action='Verify customer identity and billing account using the REQ/RITM. Check posting/rebill history before retrying.')
    for a in assets.values():
        es=byasset[a['serial']];manual=[e for e in es if e['type']=='manual']
        if manual:add('manual',a['serial'],a,'Manual adjustment needs confirmation','A manual charge or credit input exists; this package has no batch-posting confirmation.',round(sum(e['amount'] for e in manual),2),[e['source'] for e in manual],action='Confirm the manual batch imported and posted. Record the posting reference or disposition.')
        if a['state']=='In use' and not a['billed'] and not any(e['status']=='success' for e in es) and not any(e['status']=='error' for e in es) and not manual:
            assigned=day(a['assigned'])
            if assigned and (asof-assigned).days>=settings['grace_days']:
                add('unbilled',a['serial'],a,'Deployment without billing evidence',f"Assigned {assigned}; not marked billed and no success/manual entry found in supplied billing exports. Earlier billing may be outside coverage.",a['cost'],[a['source']],action='Check full billing history and eligibility before deciding whether a charge is overdue.')
        if a['serial'] in previous and previous[a['serial']]['state']=='In use' and a['state']=='In stock':
            add('return',a['serial'],a,'Returned asset needs readiness review','Observed In use → In stock between inventory exports. Inspection and reservation status are not supplied.',None,[previous[a['serial']]['source'],a['source']],action='Confirm inspection and redeployment readiness before treating it as available.')
    for s in set(dupes):add('duplicate',s,assets[s],'Duplicate inventory serial','Multiple rows have the same normalized serial in the latest snapshot.',evidence=[assets[s]['source']],priority='High',action='Correct or disambiguate the source records before using this asset in stock counts.')
    missing_rows=[r for r in inventory_rows if not serial(r['values'].get('serial_number'))]
    for r in missing_rows:
        v=r['values'];key=str(v.get('purchase_line') or v.get('purchase_line.purchase_order.number') or r['row'])
        a=dict(serial='Missing serial · '+key,model=v.get('model'),location=v.get('purchase_line.purchase_order.ship_to'),assignee=v.get('assigned_to'))
        add('missing_serial',key,a,'Inventory serial missing','An inventory record lacks a serial and cannot be linked safely to billing or counted as a tracked asset.',evidence=[r['source']],priority='High',action='Recover the serial or stable asset ID from ServiceNow; do not match by employee name alone.')
    decisions={r['id']:dict(r) for r in db.execute('SELECT * FROM decisions')}
    for issue in issues:
        d=decisions.get(issue['id'])
        if d:
            issue.update(owner=d['owner'],note=d['note'])
            issue['state']=d['state'] if d['signature']==issue['signature'] or d['state']!='Resolved' else 'Reopened'
    issues.sort(key=lambda x:(x['state']=='Resolved',x['priority']!='High',-(x['amount'] or 0)))
    groups=defaultdict(list)
    for a in assets.values():groups[(a['model'],a['location'])].append(a)
    snapshot_counts=[]
    for imp in inv:
        counts=Counter((r['values'].get('model'),r['values'].get('purchase_line.purchase_order.ship_to') or 'Unknown ship-to') for r in unpack(imp) if r['values'].get('install_status')=='In stock' and not r['values'].get('assigned_to'))
        snapshot_counts.append((imp['asof'],counts))
    stocks=[]
    for (model,loc),group in groups.items():
        qty=sum(a['state']=='In stock' and not a['assignee'] for a in group)
        deployed=sum(p['state']=='In stock' and assets[s]['state']=='In use' for s,p in previous.items() if p['model']==model and p['location']==loc and s in assets)
        rate=deployed/span if previous else None
        inbound=[]
        for o in orders:
            if o['model']!=model or o['location']!=loc or o['status']!='Ordered':continue
            received=sum(a['ticket']==o['ticket'] and a['model']==model and bool(a['received']) for a in assets.values())
            pending=max(0,o['quantity']-received)
            if pending:inbound.append(dict(o,pending=pending))
        timely=sum(o['pending'] for o in inbound if day(o['eta']) and asof<=day(o['eta'])<=asof+timedelta(days=settings['lead_days']))
        total_inbound=sum(o['pending'] for o in inbound)
        target=math.ceil(rate*(settings['lead_days']+settings['buffer_days'])) if rate is not None else None
        reorder=max(0,target-qty-timely) if target is not None else None
        cover=round(qty/rate,1) if rate else None
        risk='Review' if reorder and reorder>0 else 'Watch' if any(not day(o['eta']) or day(o['eta'])<asof for o in inbound) else 'Covered' if rate else 'No observed demand'
        history=[dict(date=snapshot_day,stock=counts[(model,loc)]) for snapshot_day,counts in snapshot_counts]
        stocks.append(dict(id=hashlib.sha256((model+'|'+loc).encode()).hexdigest()[:16],model=model,location=loc,stock=qty,deployed=deployed,span_days=span,daily_rate=round(rate,3) if rate is not None else None,days_cover=cover,inbound=total_inbound,timely_inbound=timely,orders=inbound,reorder=reorder,risk=risk,history=history,assets=[a['serial'] for a in group if a['state']=='In stock']))
    stocks.sort(key=lambda x:(x['risk']!='Review',x['risk']!='Watch',-(x['reorder'] or 0),-x['deployed']))
    depimports=[i for i in imps if i['kind']=='departures'];departures=[]
    for r in latest(imps,'departures'):
        v=r['values'];pid=str(v.get('person_id'));explicit=serial(v.get('asset_serial'));matched=[a for a in assets.values() if a['person_id']==pid or (explicit and a['serial']==explicit)]
        departures.append(dict(id=pid,name=v.get('person_name'),date=v['departure_date'],manager=v.get('manager') or 'Unassigned',days=(day(v['departure_date'])-date.today()).days,assets=[dict(serial=a['serial'],model=a['model'],state=a['state']) for a in matched],join='Explicit asset link' if explicit else 'Person ID match' if matched else 'No person ID match; inventory export needs assignee IDs',source=r['source'],demo=False))
    manifest=[dict(id=i['id'],kind=i['kind'],name=i['name'],asof=i['asof'],sheet=i['sheet'],rows=len(json.loads(i['records'])),sha=i['sha'],created=i['created']) for i in imps]
    return dict(asof=asof.isoformat(),previous_asof=prevday.isoformat(),settings=settings,assets=list(assets.values()),stocks=stocks,issues=issues,events=events,departures=departures,imports=manifest,stats=dict(assets=len(assets),stock=sum(s['stock'] for s in stocks),deployments=sum(s['deployed'] for s in stocks),open_issues=sum(i['state']!='Resolved' for i in issues),reject_amount=round(sum(i['amount'] or 0 for i in issues if i['type']=='rejected' and i['state']!='Resolved'),2),stock_reviews=sum(s['risk']=='Review' for s in stocks),interval_days=span),has_departure_feed=bool(depimports))
def timeline(db,s):
    s=serial(s);imps=imports(db);result=[];asset=None
    for imp in imps:
        if imp['kind']=='inventory':
            a,_=index_assets(unpack(imp))
            if s in a:
                asset=a[s];result.append(dict(date=imp['asof'],type='Inventory snapshot',title=f"{asset['state']} · billed flag {'Yes' if asset['billed'] else 'No'}",description=f"Assignee: {asset['assignee'] or 'Not supplied'}; assignment date: {asset['assigned'] or 'Not supplied'}",source=asset['source']))
    state=build(db)
    if asset:
        for imp in imps:
            if imp['kind'] in ('bulk_orders','custom_orders'):
                for r in unpack(imp):
                    v=r['values']
                    if v.get('purchase_order')==asset['ticket']:result.append(dict(date=r['asof'],type='Order snapshot',title=f"{v.get('status')} · {v.get('ordered_quantity')} unit(s)",description=f"{asset['ticket']} / Oracle {asset['po']} · received {v.get('purchase_order.received') or 'not supplied'}",source=r['source']))
            if imp['kind']=='po_invoices':
                for r in unpack(imp):
                    v=r['values']
                    if str(v.get('Purchase Order Number'))==asset['po'] and v.get('P2P Transaction Type Description')=='AP Invoice':result.append(dict(date=r['asof'],type='Vendor invoice snapshot',title=f"{v.get('AP Invoice Number')} · {v.get('AP Invoice Payment Status')} · ${num(v.get('AP Invoice Amount')):,.2f} allocated",description='Order-level invoice allocation; do not sum repeated snapshots or repeat for each asset.',source=r['source']))
    for e in state['events']:
        if e['serial']==s:result.append(dict(date=e['asof'],type='EAM '+e['status'] if e['type']=='eam' else 'Manual input',title=f"${e['amount']:,.2f} · {e.get('period') or ''}",description=f"{e.get('request') or ''} {e.get('ritm') or ''} · {e['message'] or 'Processed successfully'}",source=e['source']))
    return dict(asset=asset,events=sorted(result,key=lambda e:e['date']),issues=[i for i in state['issues'] if i['serial']==s])
